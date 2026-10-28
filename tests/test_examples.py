"""Comprueba los normalizadores contra los datos desordenados de examples/input/."""

import csv
from pathlib import Path

from openpyxl import load_workbook

from crm_csv_cleaner.normalize import is_blank, normalize_phone, parse_date, sort_key

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"
INVALID_PHONES = {"12345"}  # casos límite intencionados del generador
INVALID_DATES = {"2025-02-30"}


def read_csv(name: str, delimiter: str, encoding: str) -> list[dict]:
    with (EXAMPLES / "input" / name).open(newline="", encoding=encoding) as f:
        return list(csv.DictReader(f, delimiter=delimiter))


def read_xlsx(name: str) -> list[dict]:
    ws = load_workbook(EXAMPLES / "input" / name, read_only=True).active
    header, *rows = ws.iter_rows(values_only=True)
    return [dict(zip(header, row, strict=True)) for row in rows]


def input_values() -> tuple[list, list]:
    """Devuelve (teléfonos, fechas) no vacíos de los tres archivos de entrada."""
    web = read_csv("reservas_web.csv", ";", "utf-8")
    tpv = read_csv("contactos_tpv.csv", ",", "cp1252")
    eventos = read_xlsx("eventos.xlsx")
    phones = [r["Móvil"] for r in web] + [r["phone"] for r in tpv] + [r["Tlf"] for r in eventos]
    dates = (
        [r["F. reserva"] for r in web] + [r["date"] for r in tpv] + [r["Fecha"] for r in eventos]
    )
    return [p for p in phones if not is_blank(p)], [d for d in dates if not is_blank(d)]


def test_every_example_phone_normalizes():
    phones, _ = input_values()
    failed = [p for p in phones if normalize_phone(p) is None and str(p) not in INVALID_PHONES]
    assert not failed
    assert all(normalize_phone(p).startswith("+34600000") for p in phones if normalize_phone(p))


def test_every_example_date_parses():
    _, dates = input_values()
    failed = [d for d in dates if parse_date(d) is None and str(d) not in INVALID_DATES]
    assert not failed


def test_expected_output_uses_shared_sort_key():
    with (EXAMPLES / "output" / "clientes_limpios.csv").open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    keys = [sort_key(r["Last Name"], r["First Name"], r["Email"]) for r in rows]
    assert keys == sorted(keys)
