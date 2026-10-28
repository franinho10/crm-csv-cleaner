"""Falla si examples/ contiene teléfonos o emails que no sean claramente ficticios.

Reglas: teléfonos solo del bloque 600 000 xxx; emails solo @example.com / @example.org.
Se revisan todos los archivos de datos (CSV, XLSX, YAML...) excepto el código .py.
"""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"

PHONE_CANDIDATE = re.compile(r"\+?\(?\d[\d\s().+-]{7,}\d")
EMAIL_CANDIDATE = re.compile(r"[^\s@,;:<>\"']+@[^\s@,;:<>\"']*")
FAKE_PHONE = re.compile(r"600000\d{3}")
FAKE_DOMAINS = {"example.com", "example.org"}


def find_violations(text: str) -> list[str]:
    violations = []
    for match in PHONE_CANDIDATE.findall(text):
        digits = re.sub(r"\D", "", match)
        if len(digits) < 9:
            continue  # fechas, códigos cortos...
        if digits.startswith("0034"):
            digits = digits[4:]
        elif digits.startswith("34") and len(digits) == 11:
            digits = digits[2:]
        if not FAKE_PHONE.fullmatch(digits):
            violations.append(f"teléfono {match.strip()!r}")
    for match in EMAIL_CANDIDATE.findall(text):
        domain = match.rsplit("@", 1)[1].strip(".").lower()
        if domain not in FAKE_DOMAINS:
            violations.append(f"email {match.strip()!r}")
    return violations


def read_texts(path: Path) -> list[str]:
    if path.suffix == ".xlsx":
        wb = load_workbook(path, read_only=True)
        return [
            str(value)
            for ws in wb
            for row in ws.iter_rows(values_only=True)
            for value in row
            if isinstance(value, str | int) and not isinstance(value, bool)
        ]
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8").splitlines()
    except UnicodeDecodeError:
        return raw.decode("cp1252").splitlines()


DATA_FILES = sorted(p for p in EXAMPLES.rglob("*") if p.is_file() and p.suffix != ".py")


def test_examples_have_data_files():
    assert any(p.suffix == ".csv" for p in DATA_FILES)
    assert any(p.suffix == ".xlsx" for p in DATA_FILES)


@pytest.mark.parametrize("path", DATA_FILES, ids=lambda p: str(p.relative_to(EXAMPLES)))
def test_example_file_is_fictitious(path):
    violations = [v for text in read_texts(path) for v in find_violations(text)]
    assert not violations, f"{path.name}: {violations[:10]}"


def test_examples_actually_contain_contacts():
    """Evita que el test pase en vacío si cambia el formato de los datos."""
    texts = [t for p in DATA_FILES for t in read_texts(p)]
    assert sum(len(PHONE_CANDIDATE.findall(t)) for t in texts) > 20
    assert sum(len(EMAIL_CANDIDATE.findall(t)) for t in texts) > 20


@pytest.mark.parametrize(
    "text",
    ["+34 612 345 678", "0034611223344", "(+34) 699-123-456", "ana@gmail.com", "x@example.net"],
)
def test_checker_flags_real_looking_data(text):
    assert find_violations(text)


@pytest.mark.parametrize(
    "text",
    [
        "+34 600 000 123",
        "0034600000999",
        "600.000.042",
        "LUCIA@Example.COM ",
        "2025-04-03",
        "12345",
    ],
)
def test_checker_accepts_fictitious_data(text):
    assert not find_violations(text)
