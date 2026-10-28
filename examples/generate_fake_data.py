"""Genera datos de ejemplo 100 % ficticios en examples/input/ y la salida esperada en
examples/output/clientes_limpios.csv.

PRIVACIDAD
- Teléfonos: SOLO el bloque 600 000 xxx (+34 600 000 000 – +34 600 000 999). Cualquier móvil
  español con formato válido podría estar asignado a una persona real; este bloque se usa para
  que se note que son datos de prueba. No llames ni escribas a estos números.
- Emails: SOLO @example.com / @example.org (dominios reservados, RFC 2606).
- Nombres: combinaciones aleatorias de listas inventadas a mano.
tests/test_privacy.py falla si algo en examples/ se sale de estas reglas.

Uso: python examples/generate_fake_data.py
(requiere `pip install -e .`; semilla fija: la salida es reproducible)
"""

import csv
import random
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

from crm_csv_cleaner.normalize import fold, sort_key

SEED = 42
N_PEOPLE = 40
HERE = Path(__file__).resolve().parent
INPUT_DIR = HERE / "input"
OUTPUT_DIR = HERE / "output"

WEB = "reservas_web.csv"
TPV = "contactos_tpv.csv"
EVENTOS = "eventos.xlsx"

# Nombres de una sola palabra para que "Nombre completo" se pueda partir sin ambigüedad.
FIRST_NAMES = [
    "Lucía",
    "Tomás",
    "Marta",
    "Andrés",
    "Noelia",
    "Iker",
    "Carmen",
    "Rubén",
    "Ainhoa",
    "Óscar",
    "Elena",
    "Jaime",
    "Sofía",
    "Héctor",
    "Paula",
    "Íñigo",
    "Nerea",
    "Mateo",
    "Lorena",
    "Adrián",
]
LAST_NAMES = [
    "Ferrán",
    "Rivas",
    "Ortega",
    "Muñoz",
    "Calvo",
    "Serrano",
    "Peña",
    "Vidal",
    "Lozano",
    "Bravo",
    "Castaño",
    "Ibáñez",
    "Robles",
    "Gallardo",
    "Soler",
    "Prieto",
    "Marín",
    "Nieto",
    "Rey",
    "Vega",
]
DOMAINS = ("example.com", "example.org")
YEAR_START = date(2025, 1, 1)
FIXED_TIMESTAMP = (2026, 1, 1, 0, 0, 0)


@dataclass
class Person:
    first: str
    last: str
    phone: str | None  # 9 dígitos nacionales, siempre "600000xxx"
    email: str | None
    # (archivo, fecha válida o None, ¿la fila incluía el email?)
    visits: list[tuple[str, date | None, bool]] = field(default_factory=list)


def slug(text: str) -> str:
    return fold(text).replace(" ", "")


def fake_phone(suffix: int) -> str:
    assert 0 <= suffix <= 999
    return f"600000{suffix:03d}"


def random_date(rng: random.Random) -> date:
    return YEAR_START + timedelta(days=rng.randrange(300))


# --- Variantes "sucias" de cada campo --------------------------------------------------------


def messy_name(rng: random.Random, text: str) -> str:
    variant = rng.choice([str.lower, str.upper, str.title, str])(text)
    return rng.choice(["{}", " {} ", "{}  "]).format(variant.replace(" ", rng.choice([" ", "  "])))


def messy_email(rng: random.Random, email: str) -> str:
    variant = rng.choice([email, email.upper(), email.capitalize()])
    return rng.choice(["{}", " {}", "{} "]).format(variant)


def web_phone(rng: random.Random, p: str) -> str:
    return rng.choice(
        [
            f"{p[:3]} {p[3:6]} {p[6:]}",
            p,
            f"+34 {p[:3]} {p[3:5]} {p[5:7]} {p[7:]}",
            f"+34{p}",
        ]
    )


def tpv_phone(rng: random.Random, p: str) -> str:
    return rng.choice([f"0034{p}", f"(+34) {p[:3]}-{p[3:6]}-{p[6:]}", f"{p[:3]}.{p[3:6]}.{p[6:]}"])


def tpv_date(rng: random.Random, d: date) -> str:
    return rng.choice([d.isoformat(), f"{d.day}-{d.month}-{d.year}"])


# --- Generación ------------------------------------------------------------------------------


def build_people(rng: random.Random, suffixes: list[int]) -> list[Person]:
    people = []
    for i in range(N_PEOPLE):
        first = rng.choice(FIRST_NAMES)
        last = f"{rng.choice(LAST_NAMES)} {rng.choice(LAST_NAMES)}"
        email = f"{slug(first)}.{slug(last.split()[0])}{i}@{rng.choice(DOMAINS)}"
        kind = rng.random()
        if kind < 0.1:  # solo email
            people.append(Person(first, last, None, email))
        elif kind < 0.25:  # solo teléfono
            people.append(Person(first, last, fake_phone(suffixes.pop()), None))
        else:
            people.append(Person(first, last, fake_phone(suffixes.pop()), email))
    return people


def assign_visits(rng: random.Random, people: list[Person]) -> None:
    files = [WEB, TPV, EVENTOS]
    for person in people:
        n = rng.choices([1, 2, 3], weights=[60, 30, 10])[0]
        chosen = rng.sample(files, n)
        if rng.random() < 0.05:  # duplicado dentro del mismo archivo
            chosen.append(chosen[0])
        for i, fname in enumerate(chosen):
            # Sin teléfono, el email es la única clave: va siempre. Con ambos, a veces falta
            # el email (nunca en la primera aparición), y el teléfono enlaza los registros.
            with_email = person.email is not None and (
                person.phone is None or i == 0 or rng.random() > 0.3
            )
            person.visits.append((fname, random_date(rng), with_email))


def rows_for(rng: random.Random, people: list[Person]) -> dict[str, list[dict]]:
    rows: dict[str, list[dict]] = {WEB: [], TPV: [], EVENTOS: []}
    for person in people:
        for fname, d, with_email in person.visits:
            email = messy_email(rng, person.email) if with_email else ""
            p = person.phone
            if fname == WEB:
                rows[WEB].append(
                    {
                        "Nombre completo": messy_name(rng, f"{person.first} {person.last}"),
                        "Móvil": web_phone(rng, p) if p else "",
                        "Correo": email,
                        "F. reserva": d.strftime("%d/%m/%Y"),
                    }
                )
            elif fname == TPV:
                rows[TPV].append(
                    {
                        "first name": messy_name(rng, person.first),
                        "last name": messy_name(rng, person.last),
                        "phone": tpv_phone(rng, p) if p else "",
                        "email": email,
                        "date": tpv_date(rng, d),
                    }
                )
            else:
                as_text = rng.random() < 0.2
                rows[EVENTOS].append(
                    {
                        "Nombre": messy_name(rng, person.first),
                        "Apellidos": messy_name(rng, person.last),
                        "Tlf": p or "",
                        "E-mail": email,
                        "Fecha": d.strftime("%d.%m.%Y")
                        if as_text
                        else datetime(d.year, d.month, d.day),
                    }
                )
    for file_rows in rows.values():
        rng.shuffle(file_rows)
    return rows


def add_edge_cases(rows: dict[str, list[dict]], people: list[Person], suffixes: list[int]) -> int:
    """Casos límite fijos. Devuelve cuántas filas deben descartarse."""
    # Descartes esperados: fila vacía, sin teléfono ni email, teléfono inválido sin email.
    rows[WEB].append({"Nombre completo": "", "Móvil": "", "Correo": "", "F. reserva": ""})
    rows[WEB].append(
        {
            "Nombre completo": "Rubén Soler Vega",
            "Móvil": "",
            "Correo": "",
            "F. reserva": "02/06/2025",
        }
    )
    rows[TPV].append(
        {
            "first name": "Elena",
            "last name": "Prieto Rey",
            "phone": "12345",
            "email": "",
            "date": "2025-07-14",
        }
    )

    # Se conservan: email inválido (queda vacío) y fecha imposible (queda vacía).
    bad_email = Person("Carmen", "Nieto Bravo", fake_phone(suffixes.pop()), None)
    bad_email.visits.append((EVENTOS, date(2025, 3, 9), False))
    rows[EVENTOS].append(
        {
            "Nombre": "Carmen",
            "Apellidos": "Nieto Bravo",
            "Tlf": bad_email.phone,
            "E-mail": "carmen.nieto.example.com",
            "Fecha": datetime(2025, 3, 9),
        }
    )
    bad_date = Person("Iker", "Calvo Marín", fake_phone(suffixes.pop()), None)
    bad_date.visits.append((TPV, None, False))
    rows[TPV].append(
        {
            "first name": "Iker",
            "last name": "Calvo Marín",
            "phone": f"0034{bad_date.phone}",
            "email": "",
            "date": "2025-02-30",
        }
    )

    # Excel convierte a número los teléfonos sin prefijo.
    numeric = Person("Paula", "Vidal Robles", fake_phone(suffixes.pop()), None)
    numeric.visits.append((EVENTOS, date(2025, 9, 21), False))
    rows[EVENTOS].append(
        {
            "Nombre": "Paula",
            "Apellidos": "Vidal Robles",
            "Tlf": int(numeric.phone),
            "E-mail": "",
            "Fecha": datetime(2025, 9, 21),
        }
    )

    people.extend([bad_email, bad_date, numeric])
    return 3


# --- Escritura -------------------------------------------------------------------------------


def write_csv(path: Path, rows: list[dict], delimiter: str, encoding: str) -> None:
    with path.open("w", newline="", encoding=encoding) as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistentes"
    ws.append(list(rows[0]))
    for row in rows:
        ws.append(list(row.values()))
        if isinstance(row["Fecha"], datetime):
            ws.cell(ws.max_row, 5).number_format = "dd/mm/yyyy"
    wb.save(path)
    _make_reproducible(path)


def _make_reproducible(path: Path) -> None:
    """openpyxl guarda la hora actual (zip y docProps); la fijamos para que no cambie el binario."""
    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                    rb"\g<1>2026-01-01T00:00:00Z",
                    data,
                )
            info.date_time = FIXED_TIMESTAMP
            dst.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)


def expected_output(people: list[Person]) -> list[dict]:
    """Salida esperada: una fila por persona con al menos una clave válida.

    Gana el registro más reciente y los campos vacíos se completan con los anteriores; como
    todas las apariciones de una persona normalizan a los mismos datos, eso equivale a sus
    datos canónicos con la fecha válida más reciente.
    """
    out = []
    for p in people:
        dates = [d for _, d, _ in p.visits if d is not None]
        email = p.email if any(with_email for *_, with_email in p.visits) else ""
        out.append(
            {
                "First Name": p.first,
                "Last Name": p.last,
                "Email": email,
                "Phone": f"+34{p.phone}" if p.phone else "",
                "Date": max(dates).isoformat() if dates else "",
                "source_files": "|".join(sorted({f for f, _, _ in p.visits}, key=sort_key)),
                "apariciones": len(p.visits),
            }
        )
    return sorted(out, key=lambda r: sort_key(r["Last Name"], r["First Name"], r["Email"]))


def main() -> None:
    rng = random.Random(SEED)
    suffixes = rng.sample(range(1000), N_PEOPLE + 10)
    people = build_people(rng, suffixes)
    assign_visits(rng, people)
    rows = rows_for(rng, people)
    discards = add_edge_cases(rows, people, suffixes)

    INPUT_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)
    write_csv(INPUT_DIR / WEB, rows[WEB], delimiter=";", encoding="utf-8")
    write_csv(INPUT_DIR / TPV, rows[TPV], delimiter=",", encoding="cp1252")
    write_xlsx(INPUT_DIR / EVENTOS, rows[EVENTOS])
    clean = expected_output(people)
    write_csv(OUTPUT_DIR / "clientes_limpios.csv", clean, delimiter=",", encoding="utf-8")

    total = sum(len(r) for r in rows.values())
    print(f"Filas generadas: {total} ({', '.join(f'{k}: {len(v)}' for k, v in rows.items())})")
    print(f"Personas únicas esperadas: {len(clean)}")
    print(f"Duplicados esperados: {total - discards - len(clean)}")
    print(f"Descartes esperados: {discards}")


if __name__ == "__main__":
    main()
